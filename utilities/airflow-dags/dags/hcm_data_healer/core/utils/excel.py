"""Styled Excel writers for reconciliation, tracer, and push outputs."""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
HDR_FONT = Font(color="FFFFFF", bold=True)
GRN_FILL = PatternFill(start_color="D6F5D6", end_color="D6F5D6", fill_type="solid")
YEL_FILL = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
RED_FILL = PatternFill(start_color="FFD7D7", end_color="FFD7D7", fill_type="solid")
SUM_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

# Large sheets skip body fill and size columns from a sample, for speed.
FILL_MAX_ROWS = 5000
WIDTH_SAMPLE_ROWS = 200


def _style(path, fill_map=None, default_fill=None):
    fill_map = fill_map or {}
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal="center")
        body_fill = fill_map.get(ws.title, default_fill)
        if body_fill and ws.max_row <= FILL_MAX_ROWS:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.fill = body_fill
        sample_max = min(ws.max_row, WIDTH_SAMPLE_ROWS) or 1
        for col in ws.iter_cols(min_row=1, max_row=sample_max):
            width = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 4, 60)
        ws.freeze_panes = "A2"
    wb.save(path)


def write_sheets(path, sheets, fill_map=None, default_fill=None):
    """sheets: ordered dict/list of (sheet_name, DataFrame)."""
    items = sheets.items() if isinstance(sheets, dict) else sheets
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        wrote_any = False
        for name, df in items:
            if df is None:
                df = pd.DataFrame()
            df.to_excel(writer, sheet_name=str(name)[:31], index=False)
            wrote_any = True
        if not wrote_any:
            pd.DataFrame().to_excel(writer, sheet_name="EMPTY", index=False)
    _style(path, fill_map=fill_map, default_fill=default_fill or SUM_FILL)
    return path
